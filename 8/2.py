# 2. Генератор Word/Excel таблиц
import tkinter as tk
from tkinter import filedialog, messagebox
from openpyxl import Workbook
from docx import Document
import random

root = tk.Tk()
root.title("Word/Excel генератор")
root.geometry("350x250")

rows_count = tk.IntVar(value=50)
auto_fill = tk.BooleanVar(value=True)


def generate_tables():
    n = rows_count.get()

    # Создание Excel
    wb = Workbook()
    ws = wb.active
    ws.title = "Измерения"
    ws['A1'] = "Номер измерения"
    ws['B1'] = "Результат"

    # Создание Word
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Номер измерения'
    hdr_cells[1].text = 'Результат'

    for i in range(n):
        result = random.randint(0, 100) if auto_fill.get() else 50
        # Excel
        ws[f'A{i + 2}'] = i + 1
        ws[f'B{i + 2}'] = result

        # Word
        row_cells = table.add_row().cells
        row_cells[0].text = str(i + 1)
        row_cells[1].text = str(result)

    filename = filedialog.asksaveasfilename(defaultextension=".xlsx")
    if filename:
        excel_name = filename
        word_name = filename.replace('.xlsx', '.docx')

        wb.save(excel_name)
        doc.save(word_name)
        messagebox.showinfo("Готово", f"Excel: {excel_name}\nWord: {word_name}")


tk.Label(root, text="Количество строк:").pack()
tk.Entry(root, textvariable=rows_count).pack(pady=5)
tk.Checkbutton(root, text="Автозаполнение (0-100)", variable=auto_fill).pack(pady=10)
tk.Button(root, text="Создать Word+Excel", command=generate_tables,
          bg="lightblue").pack(pady=20)
root.mainloop()
